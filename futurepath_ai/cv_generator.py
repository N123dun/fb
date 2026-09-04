"""
FuturePath AI - CV Generator.

Takes structured data from the CV Builder form and produces a clean,
single-column .docx CV using python-docx. Kept deliberately simple
(no tables/columns) so it renders reliably in Word, Google Docs, and
LibreOffice alike.
"""

from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_heading(doc, text):
    h = doc.add_heading(text, level=2)
    for run in h.runs:
        run.font.size = Pt(13)
    return h


def _add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def generate_cv(data: dict) -> BytesIO:
    """
    data expects keys:
      full_name, email, phone, location,
      target_career (optional string),
      summary (optional string - auto-filled if blank),
      skills (list[str]),
      education (list[dict: degree, institution, year]),
      experience (list[dict: role, organization, duration, description]),
      projects (list[dict: name, description]) - optional
      languages (list[str]) - optional
    Returns an in-memory .docx file (BytesIO), ready to send to the client.
    """
    doc = Document()

    # --- base style ---
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    # --- header: name + contact ---
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_p.add_run(data.get("full_name", "Your Name"))
    name_run.bold = True
    name_run.font.size = Pt(22)

    contact_bits = [b for b in [data.get("email"), data.get("phone"), data.get("location")] if b]
    if contact_bits:
        contact_p = doc.add_paragraph(" | ".join(contact_bits))
        contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if data.get("target_career"):
        target_p = doc.add_paragraph()
        target_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = target_p.add_run(f"Target Role: {data['target_career']}")
        run.italic = True

    doc.add_paragraph().add_run().add_break()

    # --- summary ---
    summary = data.get("summary", "").strip()
    if not summary:
        skills_preview = ", ".join(data.get("skills", [])[:3])
        summary = (
            f"Motivated and adaptable individual"
            + (f" with a growing skill set in {skills_preview}" if skills_preview else "")
            + (f", aiming to build a career in {data.get('target_career')}" if data.get("target_career") else "")
            + "."
        )
    _add_heading(doc, "Professional Summary")
    doc.add_paragraph(summary)

    # --- skills ---
    skills = data.get("skills", [])
    if skills:
        _add_heading(doc, "Skills")
        # group into a compact comma list rather than one bullet per skill
        doc.add_paragraph(", ".join(skills))

    # --- education ---
    education = data.get("education", [])
    if education:
        _add_heading(doc, "Education")
        for edu in education:
            line = doc.add_paragraph()
            r1 = line.add_run(edu.get("degree", ""))
            r1.bold = True
            institution = edu.get("institution", "")
            year = edu.get("year", "")
            tail = "  ".join([x for x in [institution, year] if x])
            if tail:
                line.add_run(f" — {tail}")

    # --- experience ---
    experience = data.get("experience", [])
    if experience:
        _add_heading(doc, "Experience")
        for exp in experience:
            line = doc.add_paragraph()
            r1 = line.add_run(exp.get("role", ""))
            r1.bold = True
            org = exp.get("organization", "")
            duration = exp.get("duration", "")
            tail = "  ".join([x for x in [org, duration] if x])
            if tail:
                line.add_run(f" — {tail}")
            if exp.get("description"):
                _add_bullet(doc, exp["description"])

    # --- projects (optional but useful for students with no job history) ---
    projects = data.get("projects", [])
    if projects:
        _add_heading(doc, "Projects")
        for proj in projects:
            line = doc.add_paragraph()
            r1 = line.add_run(proj.get("name", ""))
            r1.bold = True
            if proj.get("description"):
                line.add_run(f" — {proj['description']}")

    # --- languages ---
    languages = data.get("languages", [])
    if languages:
        _add_heading(doc, "Languages")
        doc.add_paragraph(", ".join(languages))

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
